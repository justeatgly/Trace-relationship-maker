# -*- coding: utf-8 -*-
"""根据三个输入 Excel 生成高级别到低级别需求追溯表。"""

from __future__ import annotations

import argparse
from collections import OrderedDict
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


# 无命令行参数时使用这里的配置；也建议直接按 README 中的命令行示例传参。
HIGH_LEVEL_FILE = r"C:\Users\justeatgly\req_trace\sample_data\高级别需求.xlsx"
LOW_LEVEL_FILE = r"C:\Users\justeatgly\req_trace\sample_data\低级别需求.xlsx"
TRACE_RESULT_FILE = r"C:\Users\justeatgly\req_trace\sample_data\追溯结果文件.xlsx"
OUTPUT_FILE = r"C:\Users\justeatgly\req_trace\sample_data\待检测文件.xlsx"

COL_ID = "标识"
COL_TRACE_HIGH = "高级别需求标识"
COL_TRACE_LOW = "低级别需求标识"
COL_TRACE_VERDICT = "LLM判定"

GROUP_TITLE_HIGH = "软件高级别需求"
GROUP_TITLE_LOW = "下游低级别需求"
HEADERS = ["标识", "名称", "描述", "是否需求", "是否派生", "原理"]
OUTPUT_SHEET_NAME = "Sheet1"
DEFAULT_TEMPLATE_SHEET = "Sheet2"

AFFIRMATIVE_VERDICTS = {"是", "true", "yes", "y", "1"}
NAME_COLUMNS = ("名称", "标题")
BODY_COLUMNS = ("标题和正文", "正文", "描述")


@dataclass(frozen=True)
class Requirement:
    identifier: str
    name: str
    description: str


def cell_text(value) -> str:
    """把单元格值规范为可用于匹配的文本。"""
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_header(df: pd.DataFrame) -> pd.DataFrame:
    """去除列名首尾空白和 BOM。"""
    df.columns = [cell_text(column).lstrip("\ufeff") for column in df.columns]
    return df


def _first_existing(columns: Iterable[str], candidates: Sequence[str]) -> str | None:
    available = set(columns)
    return next((candidate for candidate in candidates if candidate in available), None)


def _first_text(row: pd.Series, candidates: Sequence[str]) -> str:
    for column in candidates:
        if column in row.index:
            value = cell_text(row[column])
            if value:
                return value
    return ""


def load_requirement_lookup(file: Path | str, label: str) -> dict[str, Requirement]:
    """读取需求文件，返回以“标识”为键的需求信息。"""
    df = normalize_header(pd.read_excel(file, dtype=object))
    if COL_ID not in df.columns:
        raise KeyError(f"{label}文件缺少“{COL_ID}”列；实际表头为：{list(df.columns)}")
    if _first_existing(df.columns, BODY_COLUMNS) is None:
        raise KeyError(
            f"{label}文件缺少正文列，至少需要以下一列：{list(BODY_COLUMNS)}；"
            f"实际表头为：{list(df.columns)}"
        )

    result: dict[str, Requirement] = {}
    duplicate_ids: set[str] = set()
    for _, row in df.iterrows():
        identifier = cell_text(row[COL_ID])
        if not identifier:
            continue
        requirement = Requirement(
            identifier=identifier,
            name=_first_text(row, NAME_COLUMNS),
            description=_first_text(row, BODY_COLUMNS),
        )
        if identifier in result:
            duplicate_ids.add(identifier)
            continue
        result[identifier] = requirement

    if duplicate_ids:
        examples = "、".join(sorted(duplicate_ids)[:10])
        raise ValueError(f"{label}文件中存在重复标识：{examples}")
    return result


def is_affirmative_verdict(value) -> bool:
    """判断 LLM 判定是否为肯定值。"""
    if isinstance(value, bool):
        return value
    return cell_text(value).lower() in AFFIRMATIVE_VERDICTS


def load_trace_pairs(file: Path | str) -> list[tuple[str, str]]:
    """读取并去重所有 LLM 判定为“是”的追溯关系。"""
    df = normalize_header(pd.read_excel(file, dtype=object))
    required = [COL_TRACE_HIGH, COL_TRACE_LOW, COL_TRACE_VERDICT]
    missing = [column for column in required if column not in df.columns]
    if missing:
        raise KeyError(f"追溯关系文件缺少列：{missing}；实际表头为：{list(df.columns)}")

    pairs: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for _, row in df.iterrows():
        if not is_affirmative_verdict(row[COL_TRACE_VERDICT]):
            continue
        pair = (cell_text(row[COL_TRACE_HIGH]), cell_text(row[COL_TRACE_LOW]))
        if not pair[0] or not pair[1] or pair in seen:
            continue
        seen.add(pair)
        pairs.append(pair)
    return pairs


def group_pairs(pairs: Iterable[tuple[str, str]]) -> list[tuple[str, list[str]]]:
    """按高级别标识稳定分组，即使同一标识在关系表中并不相邻。"""
    grouped: OrderedDict[str, list[str]] = OrderedDict()
    for high_id, low_id in pairs:
        grouped.setdefault(high_id, []).append(low_id)
    return list(grouped.items())


def _new_output_workbook() -> tuple[Workbook, object, list, float | None]:
    wb = Workbook()
    ws = wb.active
    ws.title = OUTPUT_SHEET_NAME
    ws.merge_cells("A1:F1")
    ws.merge_cells("G1:L1")

    font = Font(name="等线", size=11)
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(1, 4):
        for column in range(1, 13):
            cell = ws.cell(row=row, column=column)
            cell.font = copy(font)
            cell.alignment = copy(alignment)

    # 与附件 Sheet2 的列宽一致。
    widths = [22, 20, 38, 11, 11, 34, 22, 20, 38, 11, 11, 34]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    return wb, ws, [copy(ws.cell(3, c)._style) for c in range(1, 13)], None


def _workbook_from_template(
    template_file: Path | str, template_sheet: str
) -> tuple[Workbook, object, list, float | None]:
    """保留模板目标工作表的版式，并移除其中的样例数据。"""
    wb = load_workbook(template_file)
    if template_sheet not in wb.sheetnames:
        raise KeyError(
            f"模板中找不到工作表“{template_sheet}”；可用工作表：{wb.sheetnames}"
        )
    ws = wb[template_sheet]
    data_styles = [copy(ws.cell(3, c)._style) for c in range(1, 13)]
    data_row_height = ws.row_dimensions[3].height

    for other in list(wb.worksheets):
        if other is not ws:
            wb.remove(other)
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= 3:
            ws.unmerge_cells(str(merged_range))
    if ws.max_row >= 3:
        ws.delete_rows(3, ws.max_row - 2)
    ws.title = OUTPUT_SHEET_NAME
    return wb, ws, data_styles, data_row_height


def _prepare_output_workbook(
    template_file: Path | str | None, template_sheet: str
) -> tuple[Workbook, object, list, float | None]:
    if template_file:
        return _workbook_from_template(template_file, template_sheet)
    return _new_output_workbook()


def _write_headers(ws) -> None:
    merged_ranges = {str(item) for item in ws.merged_cells.ranges}
    if "A1:F1" not in merged_ranges:
        ws.merge_cells("A1:F1")
    if "G1:L1" not in merged_ranges:
        ws.merge_cells("G1:L1")
    ws["A1"] = GROUP_TITLE_HIGH
    ws["G1"] = GROUP_TITLE_LOW
    for column, header in enumerate(HEADERS, start=1):
        ws.cell(2, column, header)
        ws.cell(2, column + 6, header)


def _apply_row_style(ws, row: int, data_styles: Sequence, row_height: float | None) -> None:
    for column, style in enumerate(data_styles, start=1):
        ws.cell(row, column)._style = copy(style)
    if row_height is not None:
        ws.row_dimensions[row].height = row_height


def build_detection_file(
    pairs: Iterable[tuple[str, str]],
    high_lookup: dict[str, Requirement],
    low_lookup: dict[str, Requirement],
    out_file: Path | str,
    template_file: Path | str | None = None,
    template_sheet: str = DEFAULT_TEMPLATE_SHEET,
) -> tuple[int, set[str], set[str]]:
    """生成追溯 Excel，返回写入条数和未找到的高/低级别标识。"""
    wb, ws, data_styles, row_height = _prepare_output_workbook(
        template_file, template_sheet
    )
    _write_headers(ws)

    missing_high: set[str] = set()
    missing_low: set[str] = set()
    row = 3
    count = 0
    for high_id, low_ids in group_pairs(pairs):
        high = high_lookup.get(high_id)
        if high is None:
            missing_high.add(high_id)
            high = Requirement(high_id, "", "")
        group_start = row
        for index, low_id in enumerate(low_ids):
            low = low_lookup.get(low_id)
            if low is None:
                missing_low.add(low_id)
                low = Requirement(low_id, "", "")
            _apply_row_style(ws, row, data_styles, row_height)
            if index == 0:
                left = [high.identifier, high.name, high.description, True, False, None]
                for column, value in enumerate(left, start=1):
                    ws.cell(row, column, value)
            right = [low.identifier, low.name, low.description, True, False, None]
            for column, value in enumerate(right, start=7):
                ws.cell(row, column, value)
            row += 1
            count += 1

        if len(low_ids) > 1:
            for column in range(1, 7):
                ws.merge_cells(
                    start_row=group_start,
                    start_column=column,
                    end_row=row - 1,
                    end_column=column,
                )

    output_path = Path(out_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return count, missing_high, missing_low


def generate_detection_file(
    high_file: Path | str,
    low_file: Path | str,
    trace_file: Path | str,
    out_file: Path | str,
    template_file: Path | str | None = None,
    template_sheet: str = DEFAULT_TEMPLATE_SHEET,
) -> tuple[int, set[str], set[str]]:
    """完整执行读取、筛选和输出流程。"""
    for file in (high_file, low_file, trace_file):
        if not Path(file).is_file():
            raise FileNotFoundError(f"找不到文件：{file}")
    if template_file and not Path(template_file).is_file():
        raise FileNotFoundError(f"找不到模板文件：{template_file}")

    high_lookup = load_requirement_lookup(high_file, "高级别需求")
    low_lookup = load_requirement_lookup(low_file, "低级别需求")
    pairs = load_trace_pairs(trace_file)
    return build_detection_file(
        pairs,
        high_lookup,
        low_lookup,
        out_file,
        template_file=template_file,
        template_sheet=template_sheet,
    )


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="根据高/低级别需求和 LLM 追溯关系生成模板格式 Excel。"
    )
    parser.add_argument("high_file", nargs="?", help="高级别需求 Excel")
    parser.add_argument("low_file", nargs="?", help="低级别需求 Excel")
    parser.add_argument("trace_file", nargs="?", help="追溯关系 Excel")
    parser.add_argument("output_file", nargs="?", help="输出 Excel")
    parser.add_argument("--template", help="可选：版式模板 Excel")
    parser.add_argument(
        "--template-sheet",
        default=DEFAULT_TEMPLATE_SHEET,
        help=f"模板工作表名（默认：{DEFAULT_TEMPLATE_SHEET}）",
    )
    args = parser.parse_args(argv)
    supplied = [args.high_file, args.low_file, args.trace_file, args.output_file]
    if any(supplied) and not all(supplied):
        parser.error("请同时提供：高级别、低级别、追溯关系和输出 Excel 路径")
    return args


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    high_file = args.high_file or HIGH_LEVEL_FILE
    low_file = args.low_file or LOW_LEVEL_FILE
    trace_file = args.trace_file or TRACE_RESULT_FILE
    output_file = args.output_file or OUTPUT_FILE

    print(f"[读取] 高级别需求：{high_file}")
    print(f"[读取] 低级别需求：{low_file}")
    print(f"[读取] 追溯关系：{trace_file}")
    if args.template:
        print(f"[模板] {args.template}（工作表：{args.template_sheet}）")

    count, missing_high, missing_low = generate_detection_file(
        high_file,
        low_file,
        trace_file,
        output_file,
        template_file=args.template,
        template_sheet=args.template_sheet,
    )
    if missing_high:
        print(f"[警告] 高级别需求文件中未找到：{sorted(missing_high)}（内容留空）")
    if missing_low:
        print(f"[警告] 低级别需求文件中未找到：{sorted(missing_low)}（内容留空）")
    print(f"[输出] 已写入 {count} 条 LLM 肯定关系：{output_file}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
