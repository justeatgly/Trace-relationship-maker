import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook

from make_detection_file import generate_detection_file


def save_rows(path, headers, rows, sheet_name="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


class DetectionFileTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.high = self.root / "high.xlsx"
        self.low = self.root / "low.xlsx"
        self.trace = self.root / "trace.xlsx"
        self.output = self.root / "output.xlsx"

        save_rows(
            self.high,
            ["标识", "名称", "描述", "标题和正文", "是否需求"],
            [
                ["H-1", "高级标题1", "旧描述1", "高级正文1", False],
                ["H-2", "高级标题2", "旧描述2", "高级正文2", True],
            ],
        )
        save_rows(
            self.low,
            ["标识", "标题和正文", "是否需求"],
            [
                ["L-1", "低级正文1", False],
                ["L-2", "低级正文2", False],
                ["L-3", "低级正文3", True],
            ],
        )
        # H-1 故意不连续，用来验证全局分组；包含否定项和重复项。
        save_rows(
            self.trace,
            ["高级别需求标识", "低级别需求标识", "LLM判定", "判定理由"],
            [
                ["H-1", "L-1", "是", "匹配"],
                ["H-2", "L-2", "否", "不匹配"],
                ["H-1", "L-3", " 是 ", "匹配"],
                ["H-1", "L-1", "是", "重复"],
                ["H-2", "L-2", True, "匹配"],
            ],
        )

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_filters_groups_merges_and_sets_requirement_true(self):
        count, missing_high, missing_low = generate_detection_file(
            self.high, self.low, self.trace, self.output
        )
        self.assertEqual(count, 3)
        self.assertEqual(missing_high, set())
        self.assertEqual(missing_low, set())

        wb = load_workbook(self.output)
        ws = wb["Sheet1"]
        self.assertEqual(ws["A1"].value, "软件高级别需求")
        self.assertEqual(ws["G1"].value, "下游低级别需求")
        self.assertEqual(
            [ws.cell(2, c).value for c in range(1, 7)],
            ["标识", "名称", "描述", "是否需求", "是否派生", "原理"],
        )
        self.assertEqual(
            [ws.cell(3, c).value for c in (1, 2, 3, 4)],
            ["H-1", "高级标题1", "高级正文1", True],
        )
        self.assertEqual(ws["G3"].value, "L-1")
        self.assertEqual(ws["G4"].value, "L-3")
        self.assertEqual(ws["G5"].value, "L-2")
        self.assertTrue(ws["J3"].value)
        self.assertTrue(ws["J4"].value)
        self.assertTrue(ws["J5"].value)
        for column in "ABCDEF":
            self.assertIn(f"{column}3:{column}4", ws.merged_cells)

    def test_reuses_template_sheet_layout(self):
        template = self.root / "template.xlsx"
        wb = Workbook()
        wb.active.title = "其他"
        ws = wb.create_sheet("Sheet2")
        ws.merge_cells("A1:F1")
        ws.merge_cells("G1:L1")
        ws.column_dimensions["C"].width = 47
        for column in range(1, 13):
            ws.cell(2, column).value = "旧表头"
            alignment = copy(ws.cell(3, column).alignment)
            alignment.vertical = "center"
            alignment.wrap_text = True
            ws.cell(3, column).alignment = alignment
        ws["A3"] = "旧数据"
        ws.merge_cells("A3:A4")
        wb.save(template)

        generate_detection_file(
            self.high,
            self.low,
            self.trace,
            self.output,
            template_file=template,
            template_sheet="Sheet2",
        )
        result = load_workbook(self.output)
        self.assertEqual(result.sheetnames, ["Sheet1"])
        self.assertEqual(result["Sheet1"].column_dimensions["C"].width, 47)
        self.assertNotEqual(result["Sheet1"]["A3"].value, "旧数据")


if __name__ == "__main__":
    unittest.main()
